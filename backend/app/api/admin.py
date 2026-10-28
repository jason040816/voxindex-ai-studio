import os
import hashlib
import secrets
from copy import copy
from io import BytesIO
from typing import Optional
from datetime import datetime, timedelta, timezone
import time
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session
from sqlalchemy import asc, desc, func, or_
from pydantic import BaseModel, Field
from app.models import get_db
from app.models.user import User
from app.models.voice_ref import VoiceReference
from app.models.synthesis import SynthesisJob
from app.models.digital_human_model import DigitalHumanModel
from app.models.digital_human import DigitalHumanVideo
from app.models.api_access import ApiKeyUsageEvent, CustomerApiKey, IntegrationSetting, ProviderCallEvent
from app.models.credits import CreditTransaction
from app.api.voice_clones import serialize_clone
from app.services.auth import get_current_user
from app.services.credits import current_balance

router = APIRouter()

STORAGE_PATH = os.getenv("STORAGE_PATH", "../storage")
OUTPUTS_PATH = os.path.join(STORAGE_PATH, "outputs")
VOICE_REFS_PATH = os.path.join(STORAGE_PATH, "voice_refs")
DIGITAL_HUMAN_PATH = os.path.join(STORAGE_PATH, "digital_human")

# Simple in-memory cache for stats (5 min TTL)
_stats_cache = {}
_cache_ttl = 300  # 5 minutes

def get_cached_stats(key: str):
    if key in _stats_cache:
        value, timestamp = _stats_cache[key]
        if time.time() - timestamp < _cache_ttl:
            return value
    return None

def set_cached_stats(key: str, value):
    _stats_cache[key] = (value, time.time())

def clear_cached_stats(prefix: Optional[str] = None):
    if prefix is None:
        _stats_cache.clear()
        return

    for key in list(_stats_cache.keys()):
        if key.startswith(prefix):
            _stats_cache.pop(key, None)

def require_admin(current_user: User = Depends(get_current_user)) -> User:
    if not current_user.is_admin:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Admin access required")
    return current_user


def _excel_datetime(value):
    if value is not None and value.tzinfo is not None:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    return value


def get_error_code_for_job(job: SynthesisJob) -> Optional[str]:
    if job.status == "failed":
        return "TTS_FAILED"
    return None


class VoiceCloneAdminUpdate(BaseModel):
    is_enabled: bool


class UserAdminUpdate(BaseModel):
    is_banned: Optional[bool] = None
    voice_clone_limit: Optional[int] = Field(default=None, ge=0)
    synthesis_limit: Optional[int] = Field(default=None, ge=0)
    digital_human_model_limit: Optional[int] = Field(default=None, ge=0)
    digital_human_video_limit: Optional[int] = Field(default=None, ge=0)


class CustomerApiKeyCreate(BaseModel):
    user_id: int = Field(ge=1)
    name: str = Field(min_length=1, max_length=80)
    max_requests: int = Field(default=0, ge=0)


class CustomerApiKeyUpdate(BaseModel):
    max_requests: int = Field(ge=0)


class IntegrationToggle(BaseModel):
    enabled: bool


def _utc_datetime(value: Optional[datetime]) -> Optional[datetime]:
    if value is None:
        return None
    if value.tzinfo is None or value.utcoffset() is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def _api_key_payload(api_key: CustomerApiKey, owner: User, db: Session):
    usage_events = db.query(ApiKeyUsageEvent).filter(ApiKeyUsageEvent.api_key_id == api_key.id)
    clone_count = usage_events.filter(
        ApiKeyUsageEvent.operation == "voice_clone",
        ApiKeyUsageEvent.outcome == "completed",
    ).count()
    synthesis_success_count = usage_events.filter(
        ApiKeyUsageEvent.operation == "synthesis",
        ApiKeyUsageEvent.outcome == "completed",
    ).count()
    synthesis_failed_count = usage_events.filter(
        ApiKeyUsageEvent.operation == "synthesis",
        ApiKeyUsageEvent.outcome == "failed",
    ).count()
    synthesized_characters = usage_events.with_entities(
        func.coalesce(func.sum(ApiKeyUsageEvent.characters), 0),
    ).scalar() or 0
    points_used = usage_events.with_entities(
        func.coalesce(func.sum(ApiKeyUsageEvent.credits_used), 0),
    ).scalar() or 0
    return {
        "id": api_key.id,
        "username": owner.username,
        "email": owner.email,
        "name": api_key.name,
        "prefix": api_key.token_prefix,
        "request_count": api_key.request_count or 0,
        "max_requests": api_key.max_requests or 0,
        "is_exhausted": bool(api_key.max_requests and (api_key.request_count or 0) >= api_key.max_requests),
        "voice_clone_count": clone_count,
        "synthesis_count": synthesis_success_count + synthesis_failed_count,
        "synthesis_success_count": synthesis_success_count,
        "synthesis_failed_count": synthesis_failed_count,
        "synthesized_characters": synthesized_characters,
        "points_used": points_used,
        "created_at": _utc_datetime(api_key.created_at),
        "last_used_at": _utc_datetime(api_key.last_used_at),
        "revoked_at": _utc_datetime(api_key.revoked_at),
        "is_active": api_key.revoked_at is None,
    }


@router.get("/api-keys")
def list_customer_api_keys(
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
    user_id: Optional[int] = Query(None, ge=1),
):
    del admin
    query = db.query(CustomerApiKey, User).join(User, User.id == CustomerApiKey.user_id)
    if user_id is not None:
        query = query.filter(CustomerApiKey.user_id == user_id)
    rows = query.order_by(desc(CustomerApiKey.created_at), desc(CustomerApiKey.id)).all()
    return [_api_key_payload(api_key, owner, db) for api_key, owner in rows]


@router.post("/api-keys", status_code=status.HTTP_201_CREATED)
def create_customer_api_key(
    payload: CustomerApiKeyCreate,
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    clean_name = payload.name.strip()
    if not clean_name:
        raise HTTPException(status_code=422, detail="请输入 Key 名称")

    owner = db.query(User).filter(User.id == payload.user_id, User.is_admin.is_(False)).first()
    if not owner:
        raise HTTPException(status_code=404, detail="未找到客户账号")

    raw_key = f"dx_live_{secrets.token_urlsafe(32)}"
    api_key = CustomerApiKey(
        user_id=owner.id,
        name=clean_name,
        token_hash=hashlib.sha256(raw_key.encode("utf-8")).hexdigest(),
        token_prefix=raw_key[:16],
        max_requests=payload.max_requests,
        created_at=datetime.now(timezone.utc),
    )
    db.add(api_key)
    db.commit()
    db.refresh(api_key)
    return {**_api_key_payload(api_key, owner, db), "api_key": raw_key}


@router.patch("/api-keys/{api_key_id}")
def update_customer_api_key(
    api_key_id: int,
    payload: CustomerApiKeyUpdate,
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    del admin
    api_key = db.query(CustomerApiKey).filter(CustomerApiKey.id == api_key_id).first()
    if not api_key:
        raise HTTPException(status_code=404, detail="API Key not found")
    if api_key.revoked_at is not None:
        raise HTTPException(status_code=400, detail="不能修改已停用的 API Key")
    api_key.max_requests = payload.max_requests
    db.commit()
    db.refresh(api_key)
    owner = db.query(User).filter(User.id == api_key.user_id).first()
    return _api_key_payload(api_key, owner, db)


@router.get("/api-keys/{api_key_id}/usage")
def get_customer_api_key_usage(
    api_key_id: int,
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
    page: int = Query(1, ge=1),
    limit: int = Query(100, ge=1, le=200),
):
    del admin
    api_key = db.query(CustomerApiKey).filter(CustomerApiKey.id == api_key_id).first()
    if not api_key:
        raise HTTPException(status_code=404, detail="API Key not found")
    owner = db.query(User).filter(User.id == api_key.user_id).first()
    if not owner:
        raise HTTPException(status_code=404, detail="API Key owner not found")
    summary = _api_key_payload(api_key, owner, db)
    event_query = db.query(ApiKeyUsageEvent).filter(ApiKeyUsageEvent.api_key_id == api_key.id)
    total_events = event_query.count()
    events = event_query.order_by(
        desc(ApiKeyUsageEvent.created_at), desc(ApiKeyUsageEvent.id),
    ).offset((page - 1) * limit).limit(limit).all()
    return {
        "summary": summary,
        "request_events_total": total_events,
        "page": page,
        "limit": limit,
        "pages": (total_events + limit - 1) // limit if total_events else 0,
        "request_events": [
            {
                "id": event.id,
                "method": event.method,
                "path": event.path,
                "operation": event.operation,
                "status_code": event.status_code,
                "outcome": event.outcome,
                "characters": event.characters,
                "credits_used": event.credits_used,
                "created_at": _utc_datetime(event.created_at),
            }
            for event in events
        ],
    }


@router.delete("/api-keys/{api_key_id}")
def revoke_customer_api_key(
    api_key_id: int,
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    api_key = db.query(CustomerApiKey).filter(CustomerApiKey.id == api_key_id).first()
    if not api_key:
        raise HTTPException(status_code=404, detail="API Key not found")
    if api_key.revoked_at is None:
        api_key.revoked_at = datetime.now(timezone.utc)
        db.commit()
    return {"id": api_key.id, "is_active": False}


@router.get("/integrations/tencent")
def get_tencent_integration(
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    del admin
    setting = db.query(IntegrationSetting).filter(IntegrationSetting.name == "tencent_flowtts_enabled").first()
    secret_id = os.getenv("TENCENTCLOUD_SECRET_ID", "").strip()
    secret_key = os.getenv("TENCENTCLOUD_SECRET_KEY", "").strip()
    return {
        "enabled": setting.value.lower() == "true" if setting else True,
        "configured": bool(secret_id and secret_key),
        "secret_id_masked": f"••••{secret_id[-4:]}" if len(secret_id) >= 4 else ("••••" if secret_id else ""),
        "provider": os.getenv("TTS_PROVIDER", "indextts2").strip().lower(),
        "credential_source": "server_environment",
    }


@router.patch("/integrations/tencent")
def set_tencent_integration(
    payload: IntegrationToggle,
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    setting = db.query(IntegrationSetting).filter(IntegrationSetting.name == "tencent_flowtts_enabled").first()
    if setting is None:
        setting = IntegrationSetting(name="tencent_flowtts_enabled", value=str(payload.enabled).lower())
        db.add(setting)
    else:
        setting.value = str(payload.enabled).lower()
    db.commit()
    return get_tencent_integration(admin, db)


def _user_usage(user_id: int, db: Session):
    return {
        "voice_clones": db.query(VoiceReference).filter(
            VoiceReference.user_id == user_id,
            VoiceReference.is_clone.is_(True),
        ).count(),
        "audio_generations": db.query(SynthesisJob).filter(SynthesisJob.user_id == user_id).count(),
        "digital_human_models": db.query(DigitalHumanModel).filter(DigitalHumanModel.user_id == user_id).count(),
        "digital_human_videos": db.query(DigitalHumanVideo).filter(DigitalHumanVideo.user_id == user_id).count(),
    }


def _user_payload(user: User, db: Session):
    usage = _user_usage(user.id, db)
    limits = {
        "voice_clones": user.voice_clone_limit,
        "audio_generations": user.synthesis_limit,
        "digital_human_models": user.digital_human_model_limit,
        "digital_human_videos": user.digital_human_video_limit,
    }
    return {
        "id": user.id,
        "username": user.username,
        "email": user.email,
        "referral_source": user.referral_source,
        "is_admin": user.is_admin,
        "is_banned": user.is_banned,
        "credit_balance": current_balance(db, user.id),
        "created_at": user.created_at,
        "synthesis_count": usage["audio_generations"],
        "limits": limits,
        "usage": usage,
    }

@router.get("/users")
def list_users(
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
    page: int = Query(1, ge=1),
    limit: int = Query(50, ge=1, le=100),
    sort: str = Query("newest"),
    search: Optional[str] = Query(None),
    is_admin: Optional[bool] = Query(None),
):
    synthesis_counts = (
        db.query(
            SynthesisJob.user_id.label("user_id"),
            func.count(SynthesisJob.id).label("synthesis_count"),
        )
        .group_by(SynthesisJob.user_id)
        .subquery()
    )

    query = (
        db.query(User, func.coalesce(synthesis_counts.c.synthesis_count, 0).label("synthesis_count"))
        .outerjoin(synthesis_counts, synthesis_counts.c.user_id == User.id)
    )

    if search:
        pattern = f"%{search.strip().lower()}%"
        query = query.filter(
            or_(
                func.lower(User.username).like(pattern),
                func.lower(User.email).like(pattern),
                func.lower(User.referral_source).like(pattern),
            )
        )

    if is_admin is not None:
        query = query.filter(User.is_admin == is_admin)

    sort_map = {
        "newest": desc(User.created_at),
        "oldest": asc(User.created_at),
        "username_asc": asc(User.username),
        "username_desc": desc(User.username),
        "email_asc": asc(User.email),
        "email_desc": desc(User.email),
        "admin_first": desc(User.is_admin),
        "user_first": asc(User.is_admin),
    }
    query = query.order_by(sort_map.get(sort, desc(User.created_at)), desc(User.id))

    total = query.count()
    users = query.offset((page - 1) * limit).limit(limit).all()

    items = []
    for user, synthesis_count in users:
        item = _user_payload(user, db)
        item["synthesis_count"] = synthesis_count
        items.append(item)

    return {
        "items": items,
        "page": page,
        "limit": limit,
        "total": total,
        "pages": (total + limit - 1) // limit if total else 0,
    }


@router.patch("/users/{user_id}")
def update_user_controls(
    user_id: int,
    payload: UserAdminUpdate,
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    target = db.query(User).filter(User.id == user_id).first()
    if not target:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")
    if target.id == admin.id and payload.is_banned:
        raise HTTPException(status_code=400, detail="不能封禁当前管理员账号")
    for field in ("is_banned", "voice_clone_limit", "synthesis_limit", "digital_human_model_limit", "digital_human_video_limit"):
        value = getattr(payload, field)
        if value is not None:
            setattr(target, field, value)
    db.commit()
    db.refresh(target)
    return _user_payload(target, db)


@router.get("/users/{user_id}/resources")
def get_user_resources(
    user_id: int,
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    del admin
    target = db.query(User).filter(User.id == user_id).first()
    if not target:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")

    clones = db.query(VoiceReference).filter(
        VoiceReference.user_id == user_id,
        VoiceReference.is_clone.is_(True),
    ).order_by(desc(VoiceReference.created_at)).all()
    jobs = db.query(SynthesisJob).filter(
        SynthesisJob.user_id == user_id,
    ).order_by(desc(SynthesisJob.created_at)).all()
    models = db.query(DigitalHumanModel).filter(
        DigitalHumanModel.user_id == user_id,
    ).order_by(desc(DigitalHumanModel.created_at)).all()
    videos = db.query(DigitalHumanVideo).filter(
        DigitalHumanVideo.user_id == user_id,
    ).order_by(desc(DigitalHumanVideo.created_at)).all()

    public_base = (os.getenv("DIGITAL_HUMAN_PUBLIC_BASE_URL") or os.getenv("FRONTEND_APP_URL") or "").rstrip("/")
    return {
        "user": _user_payload(target, db),
        "voice_clones": [serialize_clone(item) for item in clones],
        "audio_generations": [
            {
                "id": item.id,
                "voice_ref_id": item.voice_ref_id,
                "audio_name": item.audio_name,
                "input_text": item.input_text,
                "status": item.status,
                "output_filename": item.output_filename,
                "audio_url": f"{public_base}/static/outputs/{item.output_filename}" if public_base and item.output_filename else None,
                "error_message": None if item.status == "completed" else "生成失败" if item.status == "failed" else None,
                "created_at": item.created_at,
                "ended_at": item.ended_at,
            }
            for item in jobs
        ],
        "digital_human_models": [
            {
                "id": item.id,
                "remote_model_id": item.remote_model_id,
                "name": item.name,
                "status": item.status,
                "cover_url": item.cover_url,
                "fail_reason": item.fail_reason,
                "created_at": item.created_at,
                "updated_at": item.updated_at,
            }
            for item in models
        ],
        "digital_human_videos": [
            {
                "id": item.id,
                "model_id": item.model_id,
                "synthesis_job_id": item.synthesis_job_id,
                "video_name": item.video_name,
                "status": item.status,
                "video_url": f"{public_base}/static/digital_human/{item.local_video_filename}" if public_base and item.local_video_filename else item.remote_video_url,
                "cover_url": item.cover_url,
                "fail_reason": item.fail_reason,
                "created_at": item.created_at,
                "updated_at": item.updated_at,
            }
            for item in videos
        ],
    }


@router.get("/users/{user_id}/transactions")
def get_user_credit_transactions(
    user_id: int,
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
    page: int = Query(1, ge=1),
    limit: int = Query(50, ge=1, le=100),
):
    del admin
    if not db.query(User.id).filter(User.id == user_id).first():
        raise HTTPException(status_code=404, detail="User not found")
    query = db.query(CreditTransaction).filter(CreditTransaction.user_id == user_id)
    total = query.count()
    rows = query.order_by(desc(CreditTransaction.created_at), desc(CreditTransaction.id)).offset((page - 1) * limit).limit(limit).all()
    return {
        "items": [
            {
                "id": row.id,
                "delta": row.delta,
                "kind": row.kind,
                "status": row.status,
                "description": row.description,
                "reference_type": row.reference_type,
                "reference_id": row.reference_id,
                "created_at": row.created_at,
            }
            for row in rows
        ],
        "page": page,
        "limit": limit,
        "total": total,
        "pages": (total + limit - 1) // limit if total else 0,
    }


@router.get("/users/{user_id}/provider-calls")
def get_user_provider_calls(
    user_id: int,
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
    page: int = Query(1, ge=1),
    limit: int = Query(50, ge=1, le=100),
):
    del admin
    if not db.query(User.id).filter(User.id == user_id).first():
        raise HTTPException(status_code=404, detail="User not found")
    query = db.query(ProviderCallEvent).filter(ProviderCallEvent.user_id == user_id)
    total = query.count()
    rows = query.order_by(desc(ProviderCallEvent.created_at), desc(ProviderCallEvent.id)).offset((page - 1) * limit).limit(limit).all()
    return {
        "items": [
            {
                "id": row.id,
                "provider": row.provider,
                "operation": row.operation,
                "method": row.method,
                "endpoint": row.endpoint,
                "status_code": row.status_code,
                "outcome": row.outcome,
                "duration_ms": row.duration_ms,
                "request_summary": row.request_summary,
                "response_summary": row.response_summary,
                "error_summary": row.error_summary,
                "reference_type": row.reference_type,
                "reference_id": row.reference_id,
                "created_at": row.created_at,
            }
            for row in rows
        ],
        "page": page,
        "limit": limit,
        "total": total,
        "pages": (total + limit - 1) // limit if total else 0,
    }


@router.get("/users/{user_id}/export.xlsx")
def export_user_usage_excel(
    user_id: int,
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    del admin
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    workbook = Workbook()
    info = workbook.active
    info.title = "用户信息"
    info.append(["用户ID", "用户名", "邮箱", "推荐渠道/公司", "导出时间"])
    info.append([user.id, user.username, user.email, user.referral_source or "", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")])

    transactions = workbook.create_sheet("积分明细")
    transactions.append(["记录ID", "时间", "类型", "积分变化", "状态", "说明", "关联类型", "关联ID"])
    transaction_rows = (db.query(CreditTransaction).filter(CreditTransaction.user_id == user_id)
                        .order_by(desc(CreditTransaction.created_at), desc(CreditTransaction.id)).yield_per(1000))
    for row in transaction_rows:
        transactions.append([row.id, _excel_datetime(row.created_at), row.kind, row.delta, row.status, row.description,
                             row.reference_type or "", row.reference_id or ""])

    calls = workbook.create_sheet("上游调用")
    calls.append(["记录ID", "时间", "服务商", "操作", "请求方法", "接口", "HTTP状态码", "结果", "耗时(ms)",
                  "请求摘要", "响应摘要", "错误", "关联类型", "关联ID"])
    call_rows = (db.query(ProviderCallEvent).filter(ProviderCallEvent.user_id == user_id)
                 .order_by(desc(ProviderCallEvent.created_at), desc(ProviderCallEvent.id)).yield_per(1000))
    for row in call_rows:
        calls.append([row.id, _excel_datetime(row.created_at), row.provider, row.operation, row.method, row.endpoint,
                      row.status_code, row.outcome, row.duration_ms, row.request_summary or "",
                      row.response_summary or "", row.error_summary or "", row.reference_type or "",
                      row.reference_id or ""])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1769AA")
        for column in sheet.columns:
            column_letter = column[0].column_letter
            max_length = min(max((len(str(cell.value or "")) for cell in column[:200]), default=10), 48)
            sheet.column_dimensions[column_letter].width = max(12, max_length + 2)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value
                if cell.column in (10, 11, 12):
                    alignment = copy(cell.alignment)
                    alignment.wrap_text = True
                    alignment.vertical = "top"
                    cell.alignment = alignment

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = quote(f"{user.username}_消费与上游调用明细.xlsx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.get("/audit/users")
def list_audit_users(
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
    page: int = Query(1, ge=1),
    limit: int = Query(50, ge=1, le=100),
):
    del admin
    query = db.query(User).filter(User.is_admin.is_(False))
    total = query.count()
    rows = query.order_by(desc(User.created_at), desc(User.id)).offset((page - 1) * limit).limit(limit).all()
    return {
        "items": [{"id": row.id, "username": row.username, "email": row.email,
                   "referral_source": row.referral_source, "created_at": row.created_at}
                  for row in rows],
        "page": page, "limit": limit, "total": total,
        "pages": (total + limit - 1) // limit if total else 0,
    }


@router.get("/audit/transactions")
def list_audit_transactions(
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
    page: int = Query(1, ge=1),
    limit: int = Query(50, ge=1, le=100),
):
    del admin
    query = (db.query(CreditTransaction, User.username, User.email, User.referral_source)
             .join(User, User.id == CreditTransaction.user_id)
             .filter(User.is_admin.is_(False)))
    total = query.count()
    rows = (query.order_by(desc(CreditTransaction.created_at), desc(CreditTransaction.id))
            .offset((page - 1) * limit).limit(limit).all())
    return {
        "items": [{"id": tx.id, "user_id": tx.user_id, "username": username, "email": email,
                   "referral_source": referral, "delta": tx.delta, "kind": tx.kind,
                   "status": tx.status, "description": tx.description,
                   "reference_type": tx.reference_type, "reference_id": tx.reference_id,
                   "created_at": tx.created_at}
                  for tx, username, email, referral in rows],
        "page": page, "limit": limit, "total": total,
        "pages": (total + limit - 1) // limit if total else 0,
    }


@router.get("/audit/provider-calls")
def list_audit_provider_calls(
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
    page: int = Query(1, ge=1),
    limit: int = Query(50, ge=1, le=100),
):
    del admin
    query = (db.query(ProviderCallEvent, User.username, User.email, User.referral_source)
             .outerjoin(User, User.id == ProviderCallEvent.user_id)
             .filter((User.id.is_(None)) | (User.is_admin.is_(False))))
    total = query.count()
    rows = (query.order_by(desc(ProviderCallEvent.created_at), desc(ProviderCallEvent.id))
            .offset((page - 1) * limit).limit(limit).all())
    return {
        "items": [{"id": event.id, "user_id": event.user_id, "username": username or "",
                   "email": email or "", "referral_source": referral or "",
                   "provider": event.provider, "operation": event.operation,
                   "method": event.method, "endpoint": event.endpoint,
                   "status_code": event.status_code, "outcome": event.outcome,
                   "duration_ms": event.duration_ms, "request_summary": event.request_summary,
                   "response_summary": event.response_summary, "error_summary": event.error_summary,
                   "reference_type": event.reference_type, "reference_id": event.reference_id,
                   "created_at": event.created_at}
                  for event, username, email, referral in rows],
        "page": page, "limit": limit, "total": total,
        "pages": (total + limit - 1) // limit if total else 0,
    }


@router.get("/audit/export.xlsx")
def export_all_user_audit_excel(
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    del admin
    workbook = Workbook()
    users_sheet = workbook.active
    users_sheet.title = "用户与推荐渠道"
    users_sheet.append(["用户ID", "用户名", "邮箱", "推荐渠道/公司", "注册时间"])
    for user in db.query(User).filter(User.is_admin.is_(False)).order_by(User.id).yield_per(1000):
        users_sheet.append([user.id, user.username, user.email, user.referral_source or "", _excel_datetime(user.created_at)])

    credits_sheet = workbook.create_sheet("所有用户积分流水")
    credits_sheet.append(["记录ID", "用户ID", "用户名", "邮箱", "推荐渠道/公司", "时间", "类型", "积分变化", "状态", "说明", "关联类型", "关联ID"])
    credit_rows = (db.query(CreditTransaction, User.username, User.email, User.referral_source)
                   .join(User, User.id == CreditTransaction.user_id).filter(User.is_admin.is_(False))
                   .order_by(desc(CreditTransaction.created_at), desc(CreditTransaction.id)).yield_per(1000))
    for tx, username, email, referral in credit_rows:
        credits_sheet.append([tx.id, tx.user_id, username, email, referral or "", _excel_datetime(tx.created_at),
                              tx.kind, tx.delta, tx.status, tx.description,
                              tx.reference_type or "", tx.reference_id or ""])

    calls_sheet = workbook.create_sheet("所有用户上游调用")
    calls_sheet.append(["记录ID", "用户ID", "用户名", "邮箱", "推荐渠道/公司", "时间", "服务商", "操作", "方法", "接口", "HTTP状态码", "结果", "耗时(ms)", "请求摘要", "响应摘要", "错误", "关联类型", "关联ID"])
    call_rows = (db.query(ProviderCallEvent, User.username, User.email, User.referral_source)
                 .outerjoin(User, User.id == ProviderCallEvent.user_id)
                 .filter((User.id.is_(None)) | (User.is_admin.is_(False)))
                 .order_by(desc(ProviderCallEvent.created_at), desc(ProviderCallEvent.id)).yield_per(1000))
    for event, username, email, referral in call_rows:
        calls_sheet.append([event.id, event.user_id, username or "", email or "", referral or "",
                            _excel_datetime(event.created_at), event.provider, event.operation, event.method,
                            event.endpoint, event.status_code, event.outcome, event.duration_ms,
                            event.request_summary or "", event.response_summary or "",
                            event.error_summary or "", event.reference_type or "", event.reference_id or ""])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1769AA")
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(max((len(str(cell.value or "")) for cell in column[:200]), default=10), 48)
            sheet.column_dimensions[letter].width = max(12, width + 2)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value
                if sheet is calls_sheet and cell.column in (14, 15, 16):
                    alignment = copy(cell.alignment)
                    alignment.wrap_text = True
                    alignment.vertical = "top"
                    cell.alignment = alignment

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = quote("全用户消费与上游调用明细.xlsx")
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"})

@router.get("/activity")
def list_activity(
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
    page: int = Query(1, ge=1),
    limit: int = Query(50, ge=1, le=100),
    days: int = Query(30, ge=1),
    status: str = Query("all"),
):
    cutoff_date = datetime.utcnow() - timedelta(days=days)
    
    query = db.query(SynthesisJob).filter(SynthesisJob.created_at >= cutoff_date)
    
    if status != "all":
        query = query.filter(SynthesisJob.status == status)
    
    query = query.order_by(desc(SynthesisJob.created_at))
    
    total = query.count()
    jobs = query.offset((page - 1) * limit).limit(limit).all()
    
    items = []
    for job in jobs:
        items.append({
            "id": job.id,
            "user_id": job.user_id,
            "voice_ref_id": job.voice_ref_id,
            "status": job.status,
            "error_code": get_error_code_for_job(job),
            "created_at": job.created_at,
            "duration": job.duration_seconds,
        })
    
    return {
        "items": items,
        "page": page,
        "limit": limit,
        "total": total,
        "pages": (total + limit - 1) // limit if total else 0,
    }


@router.get("/voice-clones")
def list_voice_clones(
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    rows = (
        db.query(VoiceReference, User)
        .join(User, User.id == VoiceReference.user_id)
        .filter(VoiceReference.is_clone.is_(True))
        .order_by(desc(VoiceReference.created_at))
        .all()
    )
    items = []
    for clone, owner in rows:
        item = serialize_clone(clone, include_owner=True)
        item["username"] = owner.username
        item["email"] = owner.email
        items.append(item)
    return items


@router.patch("/voice-clones/{clone_id}")
def set_voice_clone_enabled(
    clone_id: int,
    payload: VoiceCloneAdminUpdate,
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    clone = db.query(VoiceReference).filter(
        VoiceReference.id == clone_id,
        VoiceReference.is_clone.is_(True),
    ).first()
    if not clone:
        raise HTTPException(status_code=404, detail="未找到该音色")
    clone.is_enabled = payload.is_enabled
    clone.clone_status = "ready" if payload.is_enabled else "disabled"
    db.commit()
    db.refresh(clone)
    return serialize_clone(clone, include_owner=True)


@router.delete("/users/{user_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_user(user_id: int, admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    if admin.id == user_id:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="You cannot delete your own admin account")

    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")

    videos = db.query(DigitalHumanVideo).filter(DigitalHumanVideo.user_id == user_id).all()
    for video in videos:
        if video.local_video_filename:
            local_path = os.path.join(DIGITAL_HUMAN_PATH, video.local_video_filename)
            if os.path.exists(local_path):
                os.remove(local_path)
        db.delete(video)
    db.flush()

    models = db.query(DigitalHumanModel).filter(DigitalHumanModel.user_id == user_id).all()
    for model in models:
        db.delete(model)
    db.flush()

    jobs = db.query(SynthesisJob).filter(SynthesisJob.user_id == user_id).all()
    for job in jobs:
        if job.output_path and os.path.exists(job.output_path):
            os.remove(job.output_path)
        db.delete(job)
    db.flush()

    voice_refs = db.query(VoiceReference).filter(VoiceReference.user_id == user_id).all()
    for voice_ref in voice_refs:
        if voice_ref.file_path and os.path.exists(voice_ref.file_path):
            os.remove(voice_ref.file_path)
        else:
            fallback_path = os.path.join(VOICE_REFS_PATH, voice_ref.filename)
            if os.path.exists(fallback_path):
                os.remove(fallback_path)
        db.delete(voice_ref)
    db.flush()

    db.query(CustomerApiKey).filter(CustomerApiKey.user_id == user_id).delete(synchronize_session=False)
    db.flush()

    db.delete(user)
    db.commit()
    clear_cached_stats("users_growth_")
    clear_cached_stats("overview")

    return None

@router.get("/stats/overview")
def get_stats_overview(admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    cached = get_cached_stats("overview")
    if cached:
        return cached
    
    now = datetime.utcnow()
    month_ago = now - timedelta(days=30)
    month_before = month_ago - timedelta(days=30)
    
    # Active users (users who have synthesis jobs in the last 30 days)
    active_users = db.query(func.count(func.distinct(SynthesisJob.user_id))).filter(
        SynthesisJob.created_at >= month_ago
    ).scalar() or 0
    
    # Active users in previous month
    prev_active_users = db.query(func.count(func.distinct(SynthesisJob.user_id))).filter(
        SynthesisJob.created_at >= month_before,
        SynthesisJob.created_at < month_ago
    ).scalar() or 0
    
    active_users_change_percent = 0
    if prev_active_users > 0:
        active_users_change_percent = round(((active_users - prev_active_users) / prev_active_users) * 100, 1)
    
    # Total syntheses this month
    total_syntheses_month = db.query(func.count(SynthesisJob.id)).filter(
        SynthesisJob.created_at >= month_ago
    ).scalar() or 0
    
    # Average duration
    avg_duration = db.query(func.avg(SynthesisJob.duration_seconds)).filter(
        SynthesisJob.duration_seconds.isnot(None),
        SynthesisJob.created_at >= month_ago
    ).scalar() or 0
    avg_duration_seconds = round(avg_duration) if avg_duration else 0
    
    # Error rate
    failed_count = db.query(func.count(SynthesisJob.id)).filter(
        SynthesisJob.status == "failed",
        SynthesisJob.created_at >= month_ago
    ).scalar() or 0
    error_rate_percent = 0
    if total_syntheses_month > 0:
        error_rate_percent = round((failed_count / total_syntheses_month) * 100, 1)
    
    result = {
        "active_users": active_users,
        "active_users_change_percent": active_users_change_percent,
        "total_syntheses_month": total_syntheses_month,
        "avg_duration_seconds": avg_duration_seconds,
        "error_rate_percent": error_rate_percent,
    }
    
    set_cached_stats("overview", result)
    return result


@router.get("/stats/users/growth")
def get_users_growth(
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
    period: str = Query("monthly"),
):
    cache_key = f"users_growth_{period}"
    cached = get_cached_stats(cache_key)
    if cached:
        return cached
    
    now = datetime.utcnow()
    
    if period == "daily":
        # Last 7 days
        points = []
        for offset in range(6, -1, -1):
            date = (now - timedelta(days=offset)).date()
            date_start = datetime.combine(date, datetime.min.time())
            date_end = datetime.combine(date, datetime.max.time())
            
            total_users = db.query(func.count(func.distinct(User.id))).filter(
                User.created_at <= date_end
            ).scalar() or 0
            
            new_users = db.query(func.count(User.id)).filter(
                User.created_at >= date_start,
                User.created_at <= date_end
            ).scalar() or 0
            
            points.append({
                "date": date.isoformat(),
                "user_count": total_users,
                "new_users": new_users,
            })
    elif period == "yearly":
        # Last 6 years
        points = []
        for offset in range(5, -1, -1):
            year = now.year - offset
            year_start = datetime(year, 1, 1)
            year_end = datetime(year + 1, 1, 1)
            
            total_users = db.query(func.count(func.distinct(User.id))).filter(
                User.created_at < year_end
            ).scalar() or 0
            
            new_users = db.query(func.count(User.id)).filter(
                User.created_at >= year_start,
                User.created_at < year_end
            ).scalar() or 0
            
            points.append({
                "date": year_start.date().isoformat(),
                "user_count": total_users,
                "new_users": new_users,
            })
    else:  # monthly
        # Last 12 months
        points = []
        for offset in range(11, -1, -1):
            month_start = now - timedelta(days=30*(offset+1))
            month_end = now - timedelta(days=30*offset)
            
            total_users = db.query(func.count(func.distinct(User.id))).filter(
                User.created_at <= month_end
            ).scalar() or 0
            
            new_users = db.query(func.count(User.id)).filter(
                User.created_at >= month_start,
                User.created_at < month_end
            ).scalar() or 0
            
            points.append({
                "date": month_start.date().isoformat(),
                "user_count": total_users,
                "new_users": new_users,
            })
    
    set_cached_stats(cache_key, points)
    return points
